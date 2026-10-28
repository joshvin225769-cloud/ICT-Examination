"""Generate an evidence-backed PDF, editable DOCX, and Markdown internship report.

Install scripts/requirements-report.txt first. Run API/browser checks before this script.
Only set hostedUrl/liveVerified in docs/submission.json after an actual deployment check.
"""
from pathlib import Path
from datetime import datetime, timezone
import json
from html import escape

from PIL import Image as PILImage, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, Image, KeepTogether
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / 'docs'
OUT = DOCS / 'report'
OUT.mkdir(parents=True, exist_ok=True)
META = json.loads((DOCS / 'submission.json').read_text(encoding='utf-8'))
API = json.loads((DOCS / 'evidence/api-results.json').read_text(encoding='utf-8'))
BROWSER = json.loads((DOCS / 'evidence/browser-results.json').read_text(encoding='utf-8'))
TODAY = datetime.now(timezone.utc).strftime('%d %B %Y')
LIVE = META.get('hostedUrl')
if LIVE and not LIVE.startswith('https://'):
    raise ValueError('Hosted URL must use HTTPS.')
LIVE_STATUS = LIVE if LIVE else 'Pending production Atlas configuration and authenticated hosting deployment.'
STEM = 'CampusDesk_Internship_Report_Johan_Eljo'
GREEN, NAVY, MUTED = '#1c6d5c', '#18343d', '#5c6e72'


def architecture_image():
    image = PILImage.new('RGB', (1600, 430), '#f2f7f4')
    draw = ImageDraw.Draw(image)
    try:
        font = ImageFont.truetype('C:/Windows/Fonts/arial.ttf', 27)
        bold = ImageFont.truetype('C:/Windows/Fonts/arialbd.ttf', 34)
    except OSError:
        font = ImageFont.load_default(size=27)
        bold = ImageFont.load_default(size=34)
    blocks = [(65, 'REACT', 'Browser / UI', 'Forms, routes, state'), (580, 'EXPRESS + NODE', 'HTTP API / server', 'Auth, validation, controllers'), (1095, 'MONGODB', 'Atlas / database', 'Users and service requests')]
    for x, title, subtitle, detail in blocks:
        draw.rounded_rectangle((x, 115, x + 440, 310), radius=25, fill='white', outline='#d3e1d8', width=3)
        draw.text((x + 26, 145), title, fill=NAVY, font=bold)
        draw.text((x + 26, 205), subtitle, fill=GREEN, font=font)
        draw.text((x + 26, 255), detail, fill=MUTED, font=font)
    for x in (513, 1028):
        draw.line((x, 190, x + 58, 190), fill=GREEN, width=4)
        draw.polygon([(x + 58, 190), (x + 46, 182), (x + 46, 198)], fill=GREEN)
        draw.line((x, 235, x + 58, 235), fill='#74998a', width=3)
        draw.polygon([(x, 235), (x + 12, 227), (x + 12, 243)], fill='#74998a')
    draw.text((65, 35), 'REQUEST  >  VALIDATE  >  AUTHORIZE  >  PERSIST  >  RESPOND', fill=GREEN, font=font)
    draw.text((65, 355), 'The React client never connects directly to the database.', fill=MUTED, font=font)
    filename = OUT / 'architecture.png'
    image.save(filename)
    return filename


ARCH = architecture_image()


def para(text): return ('p', text)
def bullets(items): return ('bullets', items)
def table(headers, rows): return ('table', headers, rows)
def sub(title): return ('sub', title)
def image(file, caption, mobile=False): return ('image', file, caption, mobile)


pages = []
pages.append(('EXECUTIVE SUMMARY', [
    para('CampusDesk is an individual full-stack capstone developed by Johan Eljo of Sahrdaya College of Engineering and Technology for the ICT Academy of Kerala internship in Full Stack Application Development with ReactJS. It addresses the difficulty of tracking campus service issues reported through scattered messages and informal channels.'),
    para('The application provides students with a private workspace to submit service requests, track progress, refine open reports, and add follow-up comments. Administrators use a campus-wide queue to set priorities, update status, and communicate progress. The implementation uses React, Node.js, Express, MongoDB/Mongoose, and a clearly separated client/server architecture.'),
    para('Key activities included defining user stories and acceptance criteria, preparing the data model and API contract, implementing authentication and ownership checks, designing responsive client views, integrating the API, testing database persistence and authorization, and producing deployment and submission documentation.'),
    para(f'Local verification completed successfully: {API["passed"]} API integration scenarios and {BROWSER["passed"]} browser scenarios passed. Thirteen major client-side screens were captured, including mobile layouts. Automated accessibility scans reported no violations on the captured screens for the selected WCAG A/AA rules. These scans do not constitute complete accessibility certification.'),
    sub('Submission snapshot'),
    table(['Item', 'Recorded status'], [
        ['Project', META['project']], ['Student', META['student']], ['Institution', META['institution']],
        ['GitHub repository', META['repository']], ['Hosted project', LIVE_STATUS],
        ['Live verification', 'Completed' if META.get('liveVerified') else 'Pending; local checks do not establish hosted operation.'],
        ['Report generated', TODAY]
    ]),
    para('The current deliverable is a locally verified application and evidence-backed report. The hosted project link and live Atlas/HTTPS acceptance results must be added after deployment before the final submission is considered complete.')
]))
pages.append(('INTRODUCTION', [
    sub('Background and context'),
    para('Students regularly encounter infrastructure, network, library, academic-service, and hostel issues. When each report exists only as a message or a conversation, it is difficult for the reporting student to know what happened next and difficult for the service team to maintain a consistent queue. CampusDesk explores a single, traceable workflow for these requests.'),
    sub('Problem statement'),
    para('Campus service reports lack consistent tracking and visible progress when they are spread across informal channels. Students and service administrators need a shared record of the issue, its current status, and the updates leading to resolution.'),
    sub('Existing approaches and project rationale'),
    para('Email, spreadsheets, and group messaging can record issues but require manual coordination and repeated follow-ups. A service-request web application adds structured fields, ownership, a queue, and a persistent activity trail. This comparison is a design rationale, not the result of a formal survey or a measured study at the institution.'),
    sub('Technical foundation'),
    para('The MERN architecture separates responsibilities: React manages browser interaction [1], Express and Node.js handle HTTP requests and business rules [2, 3], and MongoDB stores structured documents accessed through Mongoose [4, 5]. The React application talks to the API and never directly to MongoDB.'),
    sub('Intended users'),
    bullets(['Student: registers and signs in, submits and tracks personal requests, edits or deletes own open reports, and adds comments.', 'Administrator: accesses the campus-wide queue, reviews request details, sets priority, changes status, and posts progress updates.', 'Evaluator: reviews the public source code, documented setup, report, screenshots, and eventual hosted application.'])
]))
pages.append(('OBJECTIVES', [
    para('The project objectives are expressed as functional outcomes and verifiable engineering goals.'),
    bullets(['Create a responsive React interface with clear navigation for sign-in, registration, overview, request list, request creation, and request detail.', 'Implement server-side authentication using bcrypt password hashing and expiring JWT-backed HttpOnly cookies.', 'Enforce authorization so a student cannot access another student’s requests and cannot become an administrator through public registration.', 'Implement request creation, retrieval, update, and deletion with server-side validation and appropriate HTTP status codes.', 'Provide administrator status and priority controls with a persistent activity record.', 'Connect the client through one API service, using an environment-based API URL and protected routes.', 'Support search, combined filters, and bounded pagination for growing request lists.', 'Handle loading, success, empty, and error states, including a visible retry path.', 'Verify core workflows through real API/database tests and browser journeys, and record evidence honestly.', 'Prepare reproducible setup, deployment configuration, screenshots, and internship-report documentation.']),
    sub('Learning objectives'),
    para('The implementation exercises component design, controlled forms, routing, context-based shared state, REST API design, Mongoose schema constraints, ownership-based access control, asynchronous error handling, responsive CSS, Git version control, browser automation, and technical reporting.'),
    sub('Measurable acceptance targets'),
    table(['Target', 'Evaluation approach'], [['390px and 1440px layouts', 'Browser captures and horizontal-overflow checks'], ['Private student data', 'Cross-user access tests for every protected operation'], ['Bounded results', 'Default 8 records; maximum 50 per page'], ['No secrets in client or response', 'Environment separation and password-hash response assertions'], ['Core flows persist', 'MongoDB assertions plus complete browser CRUD journey']])
]))
pages.append(('SCOPE AND DELIVERABLES', [
    sub('Included in the MVP'),
    bullets(['Student registration, login, current-session lookup, and logout.', 'Role-aware dashboards with scoped totals, recent requests, and category distribution.', 'Request fields: title, category, location, description, priority, status, owner, reference, timestamps, and embedded activity.', 'Student-owned request CRUD, with edit/delete restricted to Open status.', 'Administrator resolution controls and shared request comments.', 'Search across title/reference/location, status and category filtering, and pagination.', 'Responsive desktop/mobile screens, labelled fields, keyboard focus styles, safe errors, and retry controls.']),
    sub('Deliverables produced'),
    table(['Deliverable', 'Repository location'], [['React client', 'client/src/'], ['Express/MongoDB API', 'server/src/'], ['API integration tests', 'server/test/api.test.js'], ['Browser verification', 'scripts/browser-check.mjs'], ['Evidence records', 'docs/evidence/'], ['Major-view screenshots', 'docs/screenshots/'], ['Report PDF, DOCX, Markdown', 'docs/report/'], ['Deployment blueprint and guide', 'render.yaml; docs/DEPLOYMENT.md'], ['Project definition and contract', 'docs/PROJECT_BRIEF.md']]),
    sub('Deferred features'),
    para('Attachments, email notifications, staff assignment, password reset, email verification, institutional single sign-on, and multi-campus support are outside the MVP. The application is a capstone prototype rather than an institutionally approved operational system. Production availability, backups, and long-term support require deployment decisions.'),
    sub('Individual contribution'),
    para('This is an individual submission by Johan Eljo, covering the application design, implementation, integration, verification, and documentation described in this report.')
]))
pages.append(('METHODOLOGY', [
    sub('Incremental development approach'),
    para('The work was organized around a bounded MVP. User stories, acceptance criteria, wireframes, a data model, and an API contract were documented first. The backend and frontend were then built in separate folders and integrated through the shared HTTP contract. Local checks were run against isolated databases to avoid touching production data.'),
    image(ARCH, 'Figure 1. MERN request flow and separation of responsibilities.'),
    sub('Request flow: submitting a service issue'),
    bullets(['React validates the controlled form and sends JSON to POST /api/requests.', 'Express parses the body, checks the origin and session, validates fields, and calls the controller.', 'The controller uses the authenticated user as owner; ownership is never accepted from client input.', 'Mongoose validates and saves the request in MongoDB with its first activity event.', 'The API returns 201 and the created resource. React navigates to the persisted detail page or shows an error while retaining the form values.']),
    sub('Tools and design decisions'),
    para('Vite provides the client build pipeline, React Router provides route-level navigation, and Context manages authentication and notifications. The Express server uses central validation and error middleware, Mongoose models, thin routes, and controller functions. A same-origin production layout simplifies cookie handling and avoids unnecessary cross-domain configuration.'),
    para('Git commits capture definition, backend, frontend, verification, and documentation milestones. Node’s test runner with Supertest and a temporary MongoDB process validates the API. Playwright exercises the built client in Chromium, and axe checks selected accessibility rules. ReportLab and python-docx generate this report from one evidence-backed source.')
]))
pages.append(('PROJECT ACTIVITIES — DESIGN AND DATA', [
    sub('Planning and wireframes'),
    para('The desktop layout uses a persistent sidebar, breadcrumb header, summary cards, request table, and category panel. At narrow widths, navigation becomes horizontal, metrics use two columns, request rows become cards, and forms stack vertically. Authentication screens have a campus-oriented introduction beside a compact form.'),
    sub('Data modelling'),
    table(['Entity', 'Important fields / constraints'], [
        ['User', 'name; unique lowercase email; passwordHash excluded by default; student/admin role; timestamps'],
        ['ServiceRequest', 'unique CD reference; title 5–120; description 15–3000; location 3–120; category enum; priority enum; status enum; owner reference; activity; timestamps'],
        ['Embedded activity', 'actor display-name snapshot; created/updated/status/comment type; message; event time; maximum 200 events']
    ]),
    para('Users are referenced because they are shared across multiple requests. Activity events are embedded because they belong to one request and are displayed with it. Activity size is deliberately bounded; archiving would be needed for exceptionally long conversations. The database indexes unique email/reference, owner plus creation time, and status plus creation time.'),
    sub('Authorization and workflow rules'),
    bullets(['Public registration creates only students and rejects unexpected fields.', 'Student list, detail, statistics, edit, comment, and delete queries are scoped by the authenticated owner.', 'Administrator status routes require server-side administrator authorization.', 'Administrators cannot rewrite or delete another student’s original report.', 'Students can edit/delete only Open requests. Updates and comments append actor/date events.', 'Optimistic concurrency detects conflicting saves; deletion checks ownership, status, and document version.']),
    sub('Status lifecycle'),
    para('The main displayed progression is Open → In progress → Resolved → Closed. Administrators may deliberately move between the supported statuses, including reopening an issue. Every status/priority change is recorded. Closed is a distinct status and is not silently counted as Resolved in the dashboard card.')
]))
pages.append(('PROJECT ACTIVITIES — API AND SECURITY', [
    sub('API overview'),
    table(['Method and endpoint', 'Purpose / authorization'], [
        ['POST /api/auth/register', 'Create student; name, email, password; 201'],
        ['POST /api/auth/login', 'Sign in; email/password; 200 + session cookie'],
        ['POST /api/auth/logout', 'Clear browser session cookie; 200'],
        ['GET /api/auth/me', 'Current signed-in user; 200/401'],
        ['GET /api/requests', 'Scoped list with page, limit, search, status, category'],
        ['GET /api/requests/stats', 'Scoped totals and category distribution'],
        ['POST /api/requests', 'Create own request; 201'],
        ['GET /api/requests/:id', 'Owner or admin reads request'],
        ['PATCH /api/requests/:id', 'Owner edits complete content; Open only'],
        ['PATCH /api/requests/:id/status', 'Admin changes status and priority'],
        ['POST /api/requests/:id/comments', 'Owner/admin adds an activity comment; 201'],
        ['DELETE /api/requests/:id', 'Owner deletes Open request; 200'],
        ['GET /api/health', 'Database readiness; 200 or 503']
    ]),
    para('Success resources use {data: ...}; list results include pagination. Errors use {message, errors?}. Validation failures return 400, missing sessions 401, forbidden actions 403, unavailable records 404, conflicts 409, and rate-limit responses 429. Unexpected failures return a safe 500 message.'),
    sub('Security implementation'),
    para('Passwords are hashed with bcrypt at cost 12. JWTs expire after eight hours and are issued in HttpOnly, SameSite=Lax cookies; production adds Secure. Role information is loaded from the database rather than trusted from browser state. Server input schemas reject unexpected fields and query operators. Search input is escaped before regex construction. Mutating browser requests enforce the configured origin, and CORS allows only the configured UI origin.'),
    para('Helmet supplies security headers, authentication routes have stricter rate limits, JSON request size is bounded, and API responses use no-store caching. Secrets remain in server environment variables. Administrator bootstrap is server-side and requires a private password; those settings should be removed after initial provisioning. No real credentials appear in this report.')
]))
pages.append(('PROJECT ACTIVITIES — INTERFACE AND INTEGRATION', [
    sub('Client-side implementation'),
    para('Route-level pages cover login, registration, overview, request listing, request creation/editing, and request detail. Reusable components provide the shell, brand, table/cards, field wrappers, status badges, loading/error/empty views, and page headings. A centralized API module supplies the base URL, credentials, JSON parsing, and consistent errors.'),
    para('Forms use controlled fields, client-side checks, visible labels, character limits, disabled save buttons during requests, and inline validation feedback. Server validation remains authoritative. Shared authentication state is held in Context; an expired protected request clears client session state and redirects to sign-in.'),
    sub('Four-state screens'),
    table(['State', 'Implemented response'], [['Loading', 'Labelled spinner while requests or session are being loaded'], ['Success', 'Scoped counts, request results, metadata, and activity'], ['Empty', 'Helpful next step for new accounts; separate no-match filter message'], ['Error', 'Safe user-facing message and retry; forms retain entered values']]),
    sub('Challenges and fixes'),
    bullets(['Early light secondary text failed automated colour-contrast checks. The text and badge palettes were darkened and scans rerun on every captured screen.', 'Dependency audit detected advisories in an earlier router release. The router was upgraded to a patched release and the production build and browser journeys rerun.', 'A sign-out toast could persist into a subsequent sign-in. The unnecessary toast was removed to keep the session transition clear.', 'Local production-bundle testing and actual hosted verification were treated separately, so the report does not present local screenshots as proof of deployment.']),
    sub('Version control evidence'),
    para('The repository contains meaningful milestone commits for the brief, backend, frontend, verification, and documentation. The implementation was developed individually. A repository history supports source review, but does not by itself establish that live hosting or institutional rollout has occurred.')
]))
pages.append(('RESULTS & FINDINGS', [
    sub('Observed local verification results'),
    table(['Evidence', 'Result', 'Scope'], [
        ['API integration', f'{API["passed"]} passed', 'Express + isolated MongoDB 8.2.6'],
        ['Browser journeys', f'{BROWSER["passed"]} passed', 'Chromium; built React client + API + database'],
        ['Screenshots', str(len(BROWSER['screenshots'])), '1440px desktop; 390px mobile'],
        ['Accessibility scans', f'{len(BROWSER["accessibility"])} screens scanned', 'Selected axe WCAG A/AA checks; no reported violations'],
        ['Uncaught browser errors', str(len(BROWSER['pageErrors'])), 'Observed during scripted browser journey'],
        ['Hosted / Atlas flow', 'Pending', 'Not established by local temporary-database tests']
    ]),
    para(f'API evidence timestamp: {API["runAt"]}. Browser evidence timestamp: {BROWSER["runAt"]}. The browser run completed in approximately {BROWSER["durationSeconds"]} seconds on the development machine; this is test-suite duration, not an application response-time benchmark.'),
    sub('Functional findings'),
    bullets(['Created, edited, commented, and deleted requests persisted as expected in the isolated database.', 'Student data isolation applied consistently to listing, statistics, reading, editing, commenting, and deletion.', 'Only administrators could modify resolution status; student views reflected those changes after navigation/reload.', 'Filters and pagination returned the expected synthetic result sets, including empty matches.', 'A deliberately simulated local API outage produced an error view, and retry restored genuine database results.', 'Refreshing a protected request deep link served the built client and retrieved the persisted record.']),
    sub('Interpretation and limits'),
    para('These observations support functional correctness of the tested MVP paths. They do not quantify reduced resolution times, user satisfaction, performance under campus-scale traffic, or production reliability. The synthetic screenshots illustrate software behaviour and must not be interpreted as real campus service records.'),
    para('The following pages contain actual client-side screenshots. The hosted URL remains pending production deployment.')
]))

figures = [
    ('01-sign-in.png', 'RESULTS — SIGN-IN VIEW', 'The sign-in page provides labelled email/password fields and a registration link. Passwords are not displayed in the screenshot.'),
    ('02-registration.png', 'RESULTS — STUDENT REGISTRATION', 'Self-registration creates a student account. Administrator provisioning is intentionally handled privately on the server.'),
    ('04-student-dashboard.png', 'RESULTS — STUDENT OVERVIEW', 'The overview combines scoped request totals, recent items, and category distribution. All records shown are synthetic.'),
    ('05-request-list.png', 'RESULTS — REQUEST LIST', 'The list supports search, status/category filters, and pagination. Each row links to the detailed request record.'),
    ('07-new-request.png', 'RESULTS — NEW REQUEST FORM', 'Students provide a clear title, category, location, description, and urgency. Validation is implemented on both client and server.'),
    ('08-request-detail.png', 'RESULTS — REQUEST DETAIL', 'The detail view brings together the original description, lifecycle indicator, metadata, and activity conversation.'),
    ('09-admin-dashboard.png', 'RESULTS — ADMINISTRATOR OVERVIEW', 'The administrator sees campus-wide data in the same visual workspace, with a role-aware work queue.'),
    ('10-admin-resolution.png', 'RESULTS — RESOLUTION DESK', 'An administrator updates status/priority and adds a follow-up. Both changes are persisted and visible to the reporting student.'),
]
for number, (file, title, caption) in enumerate(figures, 2):
    pages.append((title, [para(caption), image(DOCS / 'screenshots' / file, f'Figure {number}. {caption}'), para('Capture environment: local production React bundle served by Express, backed by an isolated temporary MongoDB instance. Fictional demonstration users and requests only.')]))
pages.append(('RESULTS — MOBILE CLIENT VIEWS', [
    para('The application adapts to a 390px viewport. The sidebar becomes a compact navigation row, metrics form two columns, and request tables become touch-friendly cards. The captured viewports were checked for page-level horizontal overflow.'),
    ('mobile-pair', DOCS / 'screenshots/11-mobile-dashboard.png', DOCS / 'screenshots/12-mobile-requests.png'),
    para('Figures 10–11. Mobile overview and request list, captured at 390 × 844 pixels. These are real local screenshots, not wireframe mockups.'),
    para('The complete mobile form capture is also included as docs/screenshots/13-mobile-form.png. Automated checks are supplemented by labelled fields, focus states, semantic navigation, and reduced-motion CSS; screen-reader and wider browser coverage remain future verification work.')
]))
pages.append(('RESULTS — EMPTY AND ERROR STATES', [
    para('Empty and error states are first-class parts of the interface. A newly registered account receives a clear invitation to submit its first request. A failed list fetch shows a safe message and an explicit retry action.'),
    image(DOCS / 'screenshots/03-empty-dashboard.png', 'Figure 12. A new account with no submitted requests.'),
    image(DOCS / 'screenshots/06-retry-state.png', 'Figure 13. Simulated API failure with retry; the subsequent retry restored live local results.'),
    para('The error response was intercepted only inside the local browser test. No real hosted outage or production incident is claimed.')
]))
pages.append(('CONCLUSION', [
    para('CampusDesk demonstrates an end-to-end MERN service-request workflow with distinct student and administrator responsibilities. The completed local implementation covers authentication, ownership-protected CRUD, resolution updates, activity comments, scoped dashboards, filters, pagination, and responsive client-side views.'),
    para('The project met its locally testable functional objectives: requests persist, students cannot access another user’s data, administrators control status updates, and the client handles normal and failure paths clearly. The evidence consists of automated API/database tests, real browser interactions, accessibility scans, and actual screenshots.'),
    sub('Learning and reflection'),
    para('The project illustrates why backend authorization must remain authoritative, why consistent API contracts simplify integration, and why empty/error states need as much attention as success screens. Accessibility testing exposed issues that were not obvious from visual inspection. Versioned source, repeatable verification, and evidence-based documentation make the result easier for another developer or evaluator to inspect.'),
    sub('Remaining work before final submission'),
    bullets(['Provision production MongoDB Atlas access and an authenticated hosting deployment.', 'Run the live HTTPS, session, ownership, deep-link, and CRUD acceptance checks.', 'Add the real hosted URL to docs/submission.json and regenerate the report.', 'Keep the repository publicly accessible for evaluator review; public visibility has been confirmed.']),
    sub('Future improvements'),
    para('The next iteration could introduce email verification/password reset, attachments with controlled storage, staff assignment, notification delivery, query/index tuning, an external rate-limit store, and a documented archival policy for long activity histories. Broader device, assistive-technology, load, and recovery testing would be appropriate before real institutional use.')
]))
pages.append(('APPENDIX A — SETUP AND DEPLOYMENT', [
    sub('Repository and hosted links'),
    para(f'GitHub: {META["repository"]}'),
    para(f'Hosted application: {LIVE_STATUS}'),
    sub('Local setup'),
    bullets(['Install Node.js 22+ and run npm install at the repository root.', 'Copy server/.env.example to server/.env and configure MONGODB_URI, JWT_SECRET, and CLIENT_ORIGIN.', 'Optionally configure a private initial administrator account. Do not commit the real .env.', 'Run npm run dev and open http://localhost:5173.', 'For production-bundle testing: npm run build followed by npm start.']),
    sub('Environment variables'),
    table(['Variable', 'Purpose'], [['MONGODB_URI', 'Private Atlas/local database connection string'], ['JWT_SECRET', 'Private session-signing secret, at least 32 characters'], ['CLIENT_ORIGIN', 'Exact allowed browser origin; HTTPS in production'], ['PORT / NODE_ENV', 'Server listener and environment mode'], ['ADMIN_NAME / EMAIL / PASSWORD', 'Optional initial server-side administrator provisioning'], ['TRUST_PROXY', 'Explicit reverse-proxy hops; verify for chosen host'], ['VITE_API_URL', 'Public client API base; default /api, never a secret']]),
    sub('Hosting configuration'),
    para('render.yaml describes a Node web service that builds the React client and serves it with the API. The build command is npm ci --include=dev && npm run build and the start command is npm start. The health endpoint is /api/health. Production requires valid Atlas credentials, a host-compatible network allowlist, HTTPS, and the exact deployed origin. See docs/DEPLOYMENT.md for the complete post-deployment checklist.'),
    para('Record the deployed URL and live acceptance results in docs/submission.json, then regenerate this report.')
]))
pages.append(('APPENDIX B — VERIFICATION EVIDENCE', [
    sub('API scenarios recorded as passed'),
    bullets(API['checks']),
    para('Machine-readable source: docs/evidence/api-results.json. Tests are implemented in server/test/api.test.js. Test account passwords and signing secrets are generated at runtime and are not stored as public credentials.')
]))
pages.append(('APPENDIX C — BROWSER EVIDENCE AND REVIEW', [
    sub('Browser scenarios recorded as passed'),
    bullets(BROWSER['checks']),
    sub('Review boundaries'),
    para('Source: docs/evidence/browser-results.json. Chromium is the tested browser. Tests use synthetic data and an isolated database, with captured screens at 1440px and 390px. Automated accessibility findings cover the selected rules only. Production Atlas connectivity, live HTTPS operations, disaster recovery, and performance under load remain unverified.'),
    sub('Submission checklist'),
    table(['Item', 'Status'], [['Source code, README, environment examples', 'Prepared'], ['API/browser verification and screenshots', 'Completed locally'], ['PDF/DOCX/Markdown report', 'Generated from observed evidence'], ['Public GitHub access', 'Public visibility confirmed via unauthenticated GitHub API'], ['Hosted project URL and live flows', 'Pending'], ['Final report with verified hosted link', 'Regenerate after deployment']])
]))
pages.append(('APPENDIX D — REFERENCES', [
    para('The supplied ICTAK MERN Stack Project Guidelines and Reference Guide and the Summer Internship Report template informed the project organization and report structure. The official references below support the technology choices and implementation.'),
    bullets(['[1] React documentation — https://react.dev/', '[2] Express documentation — https://expressjs.com/', '[3] Node.js documentation — https://nodejs.org/en/docs', '[4] MongoDB Atlas documentation — https://www.mongodb.com/docs/atlas/', '[5] Mongoose documentation — https://mongoosejs.com/docs/', '[6] Git documentation — https://git-scm.com/doc', '[7] GitHub documentation — https://docs.github.com/', '[8] Express security best practices — https://expressjs.com/en/advanced/best-practice-security.html', '[9] MDN HTTP reference — https://developer.mozilla.org/en-US/docs/Web/HTTP', '[10] WCAG quick reference — https://www.w3.org/WAI/WCAG22/quickref/', '[11] Playwright documentation — https://playwright.dev/docs/intro', '[12] Render Node/Express deployment — https://render.com/docs/deploy-node-express-app']),
    sub('Project artifacts'),
    para('The repository README links the definition, deployment guide, testing record, screenshot gallery, and report. The Python generator regenerates the editable and PDF versions from the same text and evidence, reducing inconsistent counts or captions between formats.'),
    para('Prepared for Johan Eljo, Sahrdaya College of Engineering and Technology. Report structure follows the supplied ICTAK internship template.')
]))

pages.insert(0, ('CONTENTS', [
    para('Summer Internship Report — Full Stack Application Development with ReactJS'),
    table(['Section', 'Coverage'], [
        ['Executive Summary', 'Purpose, implementation, evidence, and submission status'],
        ['Introduction', 'Background, problem statement, context, and users'],
        ['Objectives', 'Functional, learning, and measurable acceptance goals'],
        ['Scope and Deliverables', 'MVP features, artifacts, and future work'],
        ['Methodology', 'Development approach, architecture, request flow, and tools'],
        ['Project Activities', 'Data design, API/security, frontend, and integration'],
        ['Results & Findings', 'Verification outcomes and desktop/mobile screenshots'],
        ['Conclusion', 'Outcomes, learning, and remaining deployment work'],
        ['Appendix A', 'Setup, environment variables, and deployment'],
        ['Appendix B', 'Recorded API verification scenarios'],
        ['Appendix C', 'Browser evidence and submission record'],
        ['Appendix D', 'Official references and project artifacts']
    ]),
    para(f'Submitted by {META["student"]}, {META["institution"]}. This report is accompanied by the editable DOCX, Markdown source, screenshots, and machine-readable verification records.')
]))


def write_markdown():
    content = ['# ICT ACADEMY OF KERALA', '## Summer Internship Report', '### Full Stack Application Development with ReactJS', f'# {META["project"]}', f'**Submitted by:** {META["student"]}', f'**Institution:** {META["institution"]}', f'**Generated:** {TODAY}', '**Status:** Local verification complete; live deployment pending.' if not META.get('liveVerified') else '**Status:** Live verification complete.', '']
    for heading, blocks in pages:
        content += [f'## {heading}', '']
        for block in blocks:
            if block[0] == 'p': content += [block[1], '']
            elif block[0] == 'sub': content += [f'### {block[1]}', '']
            elif block[0] == 'bullets': content += [f'- {item}' for item in block[1]] + ['']
            elif block[0] == 'table':
                content += ['| ' + ' | '.join(block[1]) + ' |', '| ' + ' | '.join(['---'] * len(block[1])) + ' |']
                content += ['| ' + ' | '.join(str(cell).replace('|', '/') for cell in row) + ' |' for row in block[2]] + ['']
            elif block[0] == 'image':
                rel = Path('..') / 'screenshots' / block[1].name if block[1].parent.name == 'screenshots' else Path(block[1].name)
                content += [f'![{block[2]}]({rel.as_posix()})', '']
            elif block[0] == 'mobile-pair':
                content += [f'![Mobile overview](../screenshots/{block[1].name})', f'![Mobile request list](../screenshots/{block[2].name})', '']
    (OUT / f'{STEM}.md').write_text('\n\n'.join(content).rstrip() + '\n', encoding='utf-8')


def write_docx():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(.7)
    section.bottom_margin = Inches(.65)
    section.left_margin = section.right_margin = Inches(.75)
    normal = doc.styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string('344D54')
    normal.paragraph_format.space_after = Pt(9)
    normal.paragraph_format.line_spacing = 1.17
    for name in ['Heading 1', 'Heading 2']:
        doc.styles[name].font.name = 'Calibri'
        doc.styles[name].font.color.rgb = RGBColor.from_string('1C6D5C')
    section.header.paragraphs[0].text = 'ICTAK  |  CampusDesk — Individual Capstone'
    section.header.paragraphs[0].style = doc.styles['Caption']
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('Johan Eljo  |  ')
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
    for text, size, shade in [('ICT ACADEMY OF KERALA', 27, '6FA1C7'), ('Summer Internship Report', 21, '145A86')]:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(20)
        run = p.add_run(text); run.bold = True; run.font.size = Pt(size); run.font.color.rgb = RGBColor.from_string(shade)
    doc.add_paragraph('\n\n')
    for text, size in [('Full Stack Application Development', 23), ('with ReactJS', 23), ('CampusDesk', 34), ('Campus Service Request Management System', 16)]:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text); run.bold = True; run.font.size = Pt(size); run.font.color.rgb = RGBColor.from_string('18343D')
    doc.add_paragraph('\n')
    for text in ['Submitted by', 'Johan Eljo', META['institution'], TODAY, 'Individual project | Local verification complete', 'Live deployment pending' if not META.get('liveVerified') else 'Live deployment verified']:
        p = doc.add_paragraph(text); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for heading, blocks in pages:
        doc.add_page_break(); doc.add_heading(heading, 1)
        for block in blocks:
            kind = block[0]
            if kind == 'p': doc.add_paragraph(block[1])
            elif kind == 'sub': doc.add_heading(block[1], 2)
            elif kind == 'bullets':
                for item in block[1]: doc.add_paragraph(item, style='List Bullet')
            elif kind == 'table':
                t = doc.add_table(rows=1, cols=len(block[1])); t.style = 'Light Shading Accent 1'
                for i, header in enumerate(block[1]): t.rows[0].cells[i].text = header
                for row in block[2]:
                    cells = t.add_row().cells
                    for i, value in enumerate(row): cells[i].text = str(value)
                doc.add_paragraph()
            elif kind == 'image':
                w, h = PILImage.open(block[1]).size
                maxw, maxh = 6.4, 5.35 if 'EMPTY' not in heading else 2.9
                scale = min(maxw / w, maxh / h)
                p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(str(block[1]), width=Inches(w * scale))
                doc.add_paragraph(block[2], style='Caption')
            elif kind == 'mobile-pair':
                p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(str(block[1]), width=Inches(2.35))
                p.add_run('   ')
                p.add_run().add_picture(str(block[2]), width=Inches(2.35))
    doc.save(OUT / f'{STEM}.docx')


def write_pdf():
    fontroot = Path('C:/Windows/Fonts')
    if (fontroot / 'arial.ttf').exists():
        for name, filename in [('Body', 'arial.ttf'), ('BodyBold', 'arialbd.ttf'), ('BodyItalic', 'ariali.ttf')]: pdfmetrics.registerFont(TTFont(name, str(fontroot / filename)))
    else:
        pdfmetrics.registerFont(TTFont('Body', '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'))
        pdfmetrics.registerFont(TTFont('BodyBold', '/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf'))
        pdfmetrics.registerFont(TTFont('BodyItalic', '/usr/share/fonts/truetype/dejavu/DejaVuSans-Oblique.ttf'))
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle('ReportBody', fontName='Body', fontSize=10, leading=14.2, textColor=colors.HexColor('#344d54'), spaceAfter=8))
    styles.add(ParagraphStyle('ReportBullet', fontName='Body', fontSize=10, leading=14.2, textColor=colors.HexColor('#344d54'), spaceAfter=5))
    styles.add(ParagraphStyle('ReportTitle', fontName='BodyBold', fontSize=19, leading=25, textColor=colors.HexColor(GREEN), spaceAfter=20))
    styles.add(ParagraphStyle('ReportSub', fontName='BodyBold', fontSize=11, leading=15, textColor=colors.HexColor(NAVY), spaceBefore=9, spaceAfter=7))
    styles.add(ParagraphStyle('ReportCaption', fontName='BodyItalic', fontSize=8, leading=12, textColor=colors.HexColor(MUTED), spaceBefore=8, spaceAfter=13))
    styles.add(ParagraphStyle('ReportCell', fontName='Body', fontSize=8.2, leading=12, textColor=colors.HexColor('#344d54')))
    styles.add(ParagraphStyle('ReportHeaderCell', fontName='BodyBold', fontSize=8.2, leading=12, textColor=colors.white))
    page_width, page_height = A4
    available = page_width - 104
    story = [Spacer(1, 600), PageBreak()]
    for heading, blocks in pages:
        if len(story) > 2: story.append(PageBreak())
        story.append(Paragraph(escape(heading), styles['ReportTitle']))
        for block in blocks:
            kind = block[0]
            if kind == 'p': story.append(Paragraph(escape(block[1]), styles['ReportBody']))
            elif kind == 'sub': story.append(Paragraph(escape(block[1]), styles['ReportSub']))
            elif kind == 'bullets':
                for item in block[1]:
                    story.append(Paragraph('•  ' + escape(item), styles['ReportBullet']))
            elif kind == 'table':
                rows = [[Paragraph(escape(str(cell)), styles['ReportHeaderCell']) for cell in block[1]]]
                rows += [[Paragraph(escape(str(cell)), styles['ReportCell']) for cell in row] for row in block[2]]
                if len(block[1]) == 2: widths = [available * .34, available * .66]
                else: widths = [available * .31, available * .23, available * .46]
                t = Table(rows, colWidths=widths, repeatRows=1, hAlign='LEFT')
                t.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, 0), colors.HexColor(GREEN)), ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.HexColor('#f0f5f2'), colors.white]), ('VALIGN', (0, 0), (-1, -1), 'TOP'), ('LEFTPADDING', (0, 0), (-1, -1), 10), ('RIGHTPADDING', (0, 0), (-1, -1), 10), ('TOPPADDING', (0, 0), (-1, -1), 7), ('BOTTOMPADDING', (0, 0), (-1, -1), 7), ('LINEBELOW', (0, 0), (-1, -1), .4, colors.HexColor('#dde7df'))]))
                story += [t, Spacer(1, 13)]
            elif kind == 'image':
                w, h = PILImage.open(block[1]).size
                maxh = 395 if heading.startswith('RESULTS —') else 170
                if heading == 'RESULTS — EMPTY AND ERROR STATES': maxh = 205
                scale = min(available / w, maxh / h)
                story.append(KeepTogether([Image(str(block[1]), width=w * scale, height=h * scale), Paragraph(escape(block[2]), styles['ReportCaption'])]))
            elif kind == 'mobile-pair':
                imgs = []
                for filename in block[1:]:
                    w, h = PILImage.open(filename).size
                    scale = min(available * .43 / w, 390 / h)
                    imgs.append(Image(str(filename), width=w * scale, height=h * scale))
                story += [Table([imgs], colWidths=[available / 2] * 2), Spacer(1, 13)]

    def paint(canvas, doc):
        canvas.saveState()
        if doc.page == 1:
            canvas.setFillColor(colors.HexColor('#6fa1c7')); canvas.setFont('BodyBold', 26)
            canvas.drawCentredString(page_width / 2, page_height - 90, 'ICT ACADEMY OF KERALA')
            canvas.setFillColor(colors.HexColor('#145a86')); canvas.setFont('BodyBold', 22)
            canvas.drawCentredString(page_width / 2, page_height - 133, 'Summer Internship Report')
            canvas.setFont('BodyBold', 21)
            canvas.drawCentredString(page_width / 2, page_height - 260, 'Full Stack Application Development')
            canvas.drawCentredString(page_width / 2, page_height - 292, 'with ReactJS')
            canvas.setStrokeColor(colors.HexColor('#cae2d6')); canvas.line(150, page_height - 329, page_width - 150, page_height - 329)
            canvas.setFillColor(colors.HexColor(GREEN)); canvas.setFont('BodyBold', 38)
            canvas.drawCentredString(page_width / 2, page_height - 390, 'CampusDesk')
            canvas.setFont('Body', 14); canvas.setFillColor(colors.HexColor(NAVY))
            canvas.drawCentredString(page_width / 2, page_height - 422, 'Campus Service Request Management System')
            canvas.setFont('Body', 11); canvas.drawCentredString(page_width / 2, 265, 'Submitted by')
            canvas.setFont('BodyBold', 17); canvas.drawCentredString(page_width / 2, 237, META['student'])
            canvas.setFont('BodyBold', 12); canvas.drawCentredString(page_width / 2, 203, 'Sahrdaya College of Engineering and Technology')
            canvas.setFont('Body', 11); canvas.drawCentredString(page_width / 2, 175, TODAY)
            canvas.setFillColor(colors.HexColor(MUTED)); canvas.setFont('Body', 9)
            canvas.drawCentredString(page_width / 2, 110, 'INDIVIDUAL CAPSTONE  |  LOCALLY VERIFIED IMPLEMENTATION')
            canvas.drawCentredString(page_width / 2, 92, 'Live deployment pending' if not META.get('liveVerified') else 'Live deployment verified')
        else:
            canvas.setFont('Body', 8); canvas.setFillColor(colors.HexColor(MUTED))
            canvas.drawString(52, page_height - 31, 'ICTAK  |  Summer Internship Report')
            canvas.drawRightString(page_width - 52, page_height - 31, 'CampusDesk — Johan Eljo')
            canvas.setStrokeColor(colors.HexColor('#cfe1d6')); canvas.line(52, page_height - 41, page_width - 52, page_height - 41)
            canvas.line(52, 42, page_width - 52, 42)
            canvas.drawString(52, 28, 'Sahrdaya College of Engineering and Technology')
            canvas.drawRightString(page_width - 52, 28, str(doc.page - 1))
        canvas.restoreState()
    doc = SimpleDocTemplate(str(OUT / f'{STEM}.pdf'), pagesize=A4, topMargin=65, bottomMargin=60, leftMargin=52, rightMargin=52, title='CampusDesk — Summer Internship Report', author=META['student'])
    doc.build(story, onFirstPage=paint, onLaterPages=paint)


def gallery():
    lines = ['# CampusDesk — client-side screenshot gallery', '', 'Actual local captures with fictional demonstration data. Desktop: 1440px; mobile: 390px. Hosted deployment remains pending.', '']
    for row in BROWSER['screenshots']:
        lines += [f'## {row["caption"]}', f'![{row["caption"]}](screenshots/{row["file"]})', '']
    (DOCS / 'SCREENSHOTS.md').write_text('\n\n'.join(lines).rstrip() + '\n', encoding='utf-8')


write_markdown()
write_docx()
write_pdf()
gallery()
print(f'Generated PDF, editable DOCX, Markdown, and screenshot gallery in {OUT}')
print(f'Evidence: {API["passed"]} API scenarios, {BROWSER["passed"]} browser scenarios, {len(BROWSER["screenshots"])} screenshots.')
